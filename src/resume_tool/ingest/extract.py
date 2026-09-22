"""File → plain text. One function per format, no intelligence."""

from __future__ import annotations

from pathlib import Path


class UnsupportedFormat(ValueError):
    """Raised for a file extension we cannot extract text from."""


def extract_text(path: Path | str) -> str:
    """Extract plain text from a resume file (.txt, .md, .pdf, .docx)."""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(path)

    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix == ".pdf":
        return _from_pdf(path)
    if suffix == ".docx":
        return _from_docx(path)
    raise UnsupportedFormat(
        f"cannot extract text from {suffix!r} — supported: .txt, .md, .pdf, .docx"
    )


def _from_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


BULLET = "•"


def _from_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))

    lines = []
    for p in document.paragraphs:
        text = p.text
        # Word list formatting carries the bullet as numbering metadata, not as a
        # character — so a list paragraph extracts as bare text and every
        # bullet-detecting parser downstream sees nothing. Normalize it back to a
        # literal marker. Found by running a real .docx resume through ingest and
        # recovering 0 of 17 bullets.
        if text.strip() and _is_list_paragraph(p) and not text.lstrip().startswith(BULLET):
            text = f"{BULLET} {text.lstrip()}"
        lines.append(text)

    # Tables are bad for ATS parsing, but resumes still use them — read them anyway
    # so ingest does not silently drop content.
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append("  ".join(cells))

    return "\n".join(lines)


def _is_list_paragraph(paragraph) -> bool:
    """True if Word is rendering this paragraph as a list item."""
    try:
        if "list" in (paragraph.style.name or "").lower():
            return True
    except Exception:
        pass
    # Direct numbering properties, set when a list is applied without a named style.
    try:
        return paragraph._p.find(".//{*}numPr") is not None
    except Exception:
        return False
