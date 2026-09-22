"""Structural lint for ATS parseability.

Deterministic, no LLM, no network. Checks the things that actually break
Workday / Greenhouse / Lever parsing, rather than guessing at style.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

import docx

# Headings parsers recognize. Anything else risks a section being ignored.
KNOWN_HEADINGS = {
    "summary", "objective", "profile", "about",
    "experience", "work experience", "professional experience", "employment",
    "employment history", "work history", "relevant experience",
    "education", "skills", "technical skills", "core skills",
    "projects", "certifications", "licenses", "awards", "honors", "honours",
    "publications", "languages", "volunteer", "leadership", "activities",
    "involvement", "interests", "references", "additional experience",
}

_EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.]+")
_PHONE = re.compile(r"(\+?\d[\d\-.\s()]{7,}\d)")


@dataclass
class Issue:
    level: str          # "error" | "warning"
    check: str
    detail: str

    def __str__(self) -> str:
        mark = "ERROR" if self.level == "error" else "warn "
        return f"{mark}  {self.check}: {self.detail}"


def lint_docx(path: Path | str) -> list[Issue]:
    """Return ATS structural problems. Empty list means it passed."""
    path = Path(path)
    d = docx.Document(str(path))
    body = d.element.body
    issues: list[Issue] = []

    if d.tables:
        issues.append(Issue("error", "tables",
            f"{len(d.tables)} table(s). Parsers commonly read cells out of order "
            "or drop them; use tab stops for alignment instead."))

    if body.findall(".//{*}txbxContent"):
        issues.append(Issue("error", "text boxes",
            "text box content is frequently skipped entirely"))

    if body.findall(".//{*}drawing") or body.findall(".//{*}pict"):
        issues.append(Issue("error", "images",
            "images carry no extractable text"))

    for cols in body.findall(".//{*}cols"):
        num = cols.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}num")
        if num and int(num) > 1:
            issues.append(Issue("error", "multi-column layout",
                f"{num} columns; parsers read straight across and interleave them"))

    for section in d.sections:
        if any(p.text.strip() for p in section.header.paragraphs):
            issues.append(Issue("error", "header content",
                "text in the page header is often never read — keep contact "
                "details in the body"))
            break
    for section in d.sections:
        if any(p.text.strip() for p in section.footer.paragraphs):
            issues.append(Issue("warning", "footer content",
                "text in the page footer is often never read"))
            break

    text = "\n".join(p.text for p in d.paragraphs)

    if not _EMAIL.search(text):
        issues.append(Issue("error", "email", "no email address found in the body"))
    if not _PHONE.search(text):
        issues.append(Issue("warning", "phone", "no phone number found in the body"))

    headings = _headings(d)
    if not headings:
        issues.append(Issue("error", "headings", "no section headings detected"))
    unknown = [h for h in headings if h.lower().strip() not in KNOWN_HEADINGS]
    if unknown:
        issues.append(Issue("warning", "nonstandard headings",
            f"{', '.join(unknown)} — parsers match on known section names"))

    bullets = [p.text for p in d.paragraphs if p.text.lstrip().startswith("•")]
    list_paras = [p for p in d.paragraphs
                  if "list" in (p.style.name or "").lower()
                  and not p.text.lstrip().startswith("•")]
    if list_paras and not bullets:
        issues.append(Issue("warning", "word list formatting",
            f"{len(list_paras)} list paragraph(s) with no literal bullet character; "
            "the marker is numbering metadata and will not survive text extraction"))

    return issues


def _headings(d: docx.Document) -> list[str]:
    """Lines that look like section headings: short, all caps or a Heading style.

    The first non-empty line is skipped: it is the candidate's name, which is
    conventionally set in capitals and would otherwise be reported as a
    nonstandard section.
    """
    out = []
    seen_first = False
    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        if not seen_first:
            seen_first = True
            continue
        if len(t) > 40:
            continue
        styled = (p.style.name or "").lower().startswith("heading")
        if styled or (t.isupper() and re.fullmatch(r"[A-Z][A-Z &/]+", t)):
            out.append(t)
    return out


def format_issues(issues: list[Issue]) -> str:
    if not issues:
        return "ATS lint: clean"
    return "\n".join(str(i) for i in issues)
