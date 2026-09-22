"""ATS-safe .docx writer.

Every structural choice here exists to survive a resume parser:

  - no tables. The single most common reason Workday and Greenhouse scramble a
    resume. Dates are placed with a right tab stop instead.
  - no text boxes, images, or multi-column layout — parsers skip or reorder them.
  - contact details in the body, never in a header/footer. Many parsers never
    read headers, so the candidate arrives with no email address.
  - literal "•" bullet characters rather than Word list formatting. Word lists
    carry the marker as numbering metadata, so the text extracts with no bullet
    at all; proven against a real resume where 0 of 17 bullets were recoverable.
  - standard section headings, since parsers match on known names.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.shared import Inches, Pt

from .document import ResumeDoc

BULLET = "•"

DEFAULT_FONT = "Calibri"
DEFAULT_SIZE = 9.5
NAME_SIZE = 17.0
HEADING_SIZE = 10.5
CONTACT_SIZE = 9.0
MARGIN_IN = 0.5
PAGE_WIDTH_IN = 8.5
BULLET_INDENT_IN = 0.16


def write_docx(doc: ResumeDoc, path: Path | str, *, font: str = DEFAULT_FONT,
               size: float = DEFAULT_SIZE, margin: float = MARGIN_IN) -> Path:
    """Render a ResumeDoc to an ATS-safe .docx and return the path."""
    path = Path(path)
    d = Document()
    _setup(d, font, size, margin)
    tab = Inches(PAGE_WIDTH_IN - 2 * margin)

    p = _para(d, space_after=1)
    _run(p, doc.contact.name.upper(), bold=True, size=NAME_SIZE)
    detail = doc.contact.detail_line()
    if detail:
        p = _para(d, space_after=3)
        _run(p, detail, size=CONTACT_SIZE)

    if doc.summary:
        _heading(d, "SUMMARY")
        _run(_para(d), doc.summary, size=size)

    for section in doc.sections:
        _heading(d, section.heading.upper())

        for i, entry in enumerate(section.entries):
            p = _para(d, space_before=(5 if i else 1), tab=tab)
            _run(p, entry.title, bold=True, size=size)
            if entry.date_range:
                _run(p, "\t" + entry.date_range, size=size)

            if entry.organization or entry.location:
                p = _para(d, tab=tab)
                _run(p, entry.organization, italic=True, size=size)
                if entry.location:
                    _run(p, "\t" + entry.location, italic=True, size=size)

            for bullet in entry.bullets:
                _bullet(d, bullet, size)

        for line in section.lines:
            p = _para(d, space_before=1)
            # "Label: values" keeps the label bold without needing a table.
            if ": " in line:
                label, rest = line.split(": ", 1)
                _run(p, label + ": ", bold=True, size=size)
                _run(p, rest, size=size)
            else:
                _run(p, line, size=size)

    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(path))
    return path


# ── internals ───────────────────────────────────────────────────────────────


def _setup(d: Document, font: str, size: float, margin: float) -> None:
    s = d.sections[0]
    s.top_margin = s.bottom_margin = Inches(margin)
    s.left_margin = s.right_margin = Inches(margin)

    normal = d.styles["Normal"]
    normal.font.name = font
    normal.font.size = Pt(size)
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0


def _para(d: Document, space_before: float = 0, space_after: float = 0, tab=None):
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if tab is not None:
        p.paragraph_format.tab_stops.add_tab_stop(tab, WD_TAB_ALIGNMENT.RIGHT)
    return p


def _run(p, text: str, *, bold: bool = False, italic: bool = False,
         size: float = DEFAULT_SIZE):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    return r


def _heading(d: Document, text: str):
    p = _para(d, space_before=7, space_after=2)
    _run(p, text, bold=True, size=HEADING_SIZE)
    return p


def _bullet(d: Document, text: str, size: float):
    p = _para(d, space_before=1)
    p.paragraph_format.left_indent = Inches(BULLET_INDENT_IN)
    p.paragraph_format.first_line_indent = Inches(-BULLET_INDENT_IN)
    _run(p, f"{BULLET}  {text}", size=size)
    return p
